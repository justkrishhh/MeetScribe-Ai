import streamlit as st
from main import generate_documentation
from docx import Document
import re
import io


st.set_page_config(
    page_title="MeetScribe Ai",
    layout="wide"
)

st.title("MeetScribe Ai")
st.write("Turn meeting transcripts into professional business documentation.")


# ------------------------------------------
# DOCUMENT TYPE SELECTION
# ------------------------------------------

document_types = [
    "Meeting Notes / SOP",
    "Minutes of Meeting (MoM)",
    "Agenda",
    "Action Items",
    "Daily Progress",
    "Questions & Follow-ups"
]

selected_documents = st.multiselect(
    "Select Document Types",
    document_types,
    default=document_types,
    placeholder="Select the documents you want to generate..."
)


# ------------------------------------------
# TRANSCRIPT
# ------------------------------------------

transcript = st.text_area(
    "Meeting Transcript",
    height=400,
    placeholder="Paste your meeting transcript here..."
)


# ------------------------------------------
# GENERATE
# ------------------------------------------

if st.button(
    "Generate",
    type="primary"
):

    if not selected_documents:
        st.warning("Please select at least one document type.")

    elif not transcript.strip():
        st.warning("Please paste a meeting transcript.")

    else:

        with st.spinner("Generating documentation..."):

            try:

                document_type = "\n".join(
                    f"- {doc}"
                    for doc in selected_documents
                )

                result = generate_documentation(
                    transcript,
                    document_type
                )

                st.session_state["result"] = result

            except Exception as e:

                st.error(f"Error: {e}")


# ------------------------------------------
# WORD DOCUMENT CREATOR
# ------------------------------------------

def create_word_document(text):

    document = Document()

    lines = text.split("\n")

    for line in lines:

        paragraph = document.add_paragraph()

        # Find **bold text**
        parts = re.split(r"(\*\*.*?\*\*)", line)

        for part in parts:

            if part.startswith("**") and part.endswith("**"):

                bold_text = part[2:-2]

                run = paragraph.add_run(bold_text)
                run.bold = True

            else:

                paragraph.add_run(part)

    file_buffer = io.BytesIO()

    document.save(file_buffer)

    file_buffer.seek(0)

    return file_buffer


# ------------------------------------------
# OUTPUT
# ------------------------------------------

if "result" in st.session_state:

    st.subheader("Generated Documentation")

    # Copyable text
    st.text_area(
        "Output",
        value=st.session_state["result"],
        height=700
    )

    # Create Word document
    word_file = create_word_document(
        st.session_state["result"]
    )

    st.download_button(
        "Download Microsoft Word Document",
        data=word_file,
        file_name="consultant_documentation.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )